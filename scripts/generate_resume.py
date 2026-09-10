from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "resume"
DOCX_PATH = OUTPUT_DIR / "Zaid_Salman_AI_Platform_Engineer_Resume_Final.docx"
PDF_PATH = OUTPUT_DIR / "Zaid_Salman_AI_Platform_Engineer_Resume_Final.pdf"

NAVY = "000000"
BLUE = "000000"
DARK = "000000"
BODY_FONT = "Calibri"
PDF_BODY_FONT = "Calibri"
PDF_BOLD_FONT = "Calibri-Bold"


SUMMARY = (
    "AI Platform Engineer with 3+ years of cloud and DevOps experience and hands-on delivery of "
    "production-oriented RAG and agentic AI systems. Builds LangChain/LangGraph workflows, FastAPI "
    "microservices, context and evaluation harnesses, human-in-the-loop agent loops, Langfuse "
    "observability, and secure AKS deployments for enterprise EU and MENA use cases."
)

SKILLS = [
    (
        "Generative AI & Agentic Systems",
        "LLMs, RAG, LangChain, LangGraph, Multi-Agent Workflows, Prompt Engineering, Tool Calling, "
        "Structured Outputs, Embeddings, Grounding, Human-in-the-Loop, Guardrails",
    ),
    (
        "Context, Harness & Loop Engineering",
        "Context Assembly, Retrieval Grounding, State Management, Schema Validation, Evals, RAGAS, "
        "Langfuse Tracing, Failure Recovery, Feedback Loops, Token/Latency/Cost Monitoring",
    ),
    (
        "AI Backend & Retrieval",
        "Python, FastAPI, Pydantic, Qdrant, Hybrid Search, Reranking, PostgreSQL, Redis, "
        "Azure OpenAI, Azure AI Foundry Integration, REST APIs, Microservices",
    ),
    (
        "AI Platform, Cloud & Governance",
        "Azure, AWS, AKS, Kubernetes, Docker, Helm, Argo CD, Terraform, GitHub Actions, CI/CD, "
        "Prometheus, Grafana, Key Vault, Managed Identities, JWT/RBAC, Consent Tracking, GDPR, PDPL, PII Redaction",
    ),
]

EXPERIENCE = [
    {
        "title": "DevOps Engineer",
        "company": "Synopsys Pvt Ltd (Client) — OPT IT Technologies (Payroll)",
        "location": "Bengaluru, India",
        "dates": "10/2024–Present",
        "bullets": [
            "Designed secure AWS and Azure environments with Terraform, sustaining 99.9% uptime while reducing infrastructure setup time by 30%.",
            "Operated Kubernetes clusters and 15+ containerized microservices, reducing cloud costs by 20% and establishing scalable foundations for AI API workloads.",
            "Engineered CI/CD deployment harnesses with automated quality gates, container promotion, health validation, and rollback controls, cutting release time by 20%.",
            "Applied production platform patterns to independent GenAI systems, deploying FastAPI, LangGraph, and RAG services on AKS with Helm, Argo CD, Terraform, and GitHub Actions.",
        ],
    },
    {
        "title": "Trainee DevOps Engineer",
        "company": "Yara Digital Farming Pvt Ltd",
        "location": "Bengaluru, India",
        "dates": "01/2023–09/2024",
        "bullets": [
            "Built GitHub Actions, Azure Pipelines, and Argo CD workflows for Python and container services, improving release velocity by 20%.",
            "Integrated AKS with Key Vault and Managed Identities, creating passwordless security patterns applicable to LLM endpoints and AI microservices.",
            "Automated Azure provisioning with Python and the Azure SDK, enabling repeatable cloud environments for data and AI application delivery.",
            "Managed 100+ Windows/Linux servers with Ansible and implemented autoscaling Jenkins workers on Azure VM Scale Sets.",
        ],
    },
]

PROJECTS = [
    {
        "title": "Enterprise Multilingual RAG & AI Platform",
        "dates": "2026",
        "tech": (
            "Python | FastAPI | LangChain | LangGraph | Azure OpenAI | Qdrant | PostgreSQL | Redis | "
            "MinIO | Langfuse | RAGAS | Docker | Kubernetes | Helm | Argo CD | Terraform"
        ),
        "bullets": [
            "Architected a 12-service FastAPI RAG platform in 16 containers spanning OCR ingestion, chunking, embeddings, Qdrant retrieval, Azure OpenAI generation, citations, and feedback.",
            "Engineered the context layer with locale-aware chunking, metadata filters, multilingual retrieval, and 70% dense / 30% keyword search across English, Arabic, French, and German.",
            "Built a LangGraph loop and LLM harness with routing, structured context, fallbacks, schema checks, Langfuse/RAGAS evaluation, grounded citations, and failure tracing.",
            "Productionized with Kubernetes, Helm, Argo CD, Terraform, and Prometheus/Grafana; enforced JWT/RBAC, PII redaction, consent tracking, GDPR erasure, and PDPL region tagging.",
            "Integrated scheduled SharePoint and Confluence synchronization with metadata mapping, incremental updates, and deletion propagation across object storage and vector indexes.",
        ],
    },
    {
        "title": "Agentic Cloud Reliability & FinOps Platform (ACRFP)",
        "dates": "2026",
        "tech": (
            "Python | FastAPI | LangChain | LangGraph | Pydantic | Azure | AKS | Terraform | Docker | "
            "Kubernetes | Helm | Argo CD | GitHub Actions"
        ),
        "bullets": [
            "Built a LangGraph agent loop with 3 specialists, 6 state-machine nodes, context-rich incident state, and 12 typed remediation tools for reliability and FinOps workflows.",
            "Engineered a safety harness with deterministic policy-as-code, structured outputs, sovereign-region controls, CRITICAL deny-by-default, HITL gates, and dual approval for HIGH-risk actions.",
            "Delivered 3 FastAPI services and 15 endpoints with audit tracing, dry-run execution, Terraform-provisioned AKS, Helm/Argo CD assets, and 11 CI tests.",
            "Implemented an approval console, webhook notifications, JSON/CSV audit exports, and a dry-run execution boundary preventing agents from directly invoking kubectl or Azure ARM.",
        ],
    },
]

CERTIFICATIONS = [
    "AWS Certified Solutions Architect – Associate (2024)",
    "Microsoft Certified: Azure Administrator Associate (AZ-104)",
    "Microsoft Certified: DevOps Engineer Expert (AZ-400)",
]


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_docx(document):
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.05

    if "Resume Heading" not in document.styles:
        heading = document.styles.add_style("Resume Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        heading = document.styles["Resume Heading"]
    heading.font.name = BODY_FONT
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(NAVY)
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.keep_with_next = True

    if "Resume Bullet" not in document.styles:
        bullet = document.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet = document.styles["Resume Bullet"]
    bullet.base_style = document.styles["Normal"]
    bullet.paragraph_format.left_indent = Inches(0.15)
    bullet.paragraph_format.first_line_indent = Inches(-0.12)
    bullet.paragraph_format.space_after = Pt(1.4)
    bullet.paragraph_format.keep_together = True


def add_docx_heading(document, text):
    p = document.add_paragraph(style="Resume Heading")
    p.add_run(text.upper())
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "7")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    p_pr.append(borders)


def add_docx_bullet(document, text):
    p = document.add_paragraph(style="Resume Bullet")
    p.add_run("• ").bold = True
    p.add_run(text)


def add_docx_role(document, title, company, location, dates):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(1.5)
    table.rows[0].cells[0].width = Inches(5.7)
    table.rows[0].cells[1].width = Inches(1.5)
    table.rows[0].cells[0].vertical_alignment = 1
    p = table.rows[0].cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(f" | {company} | {location}")
    right = table.rows[0].cells[1].paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    right.add_run(dates).bold = True
    for cell in table.rows[0].cells:
        cell.margin_top = 0
        cell.margin_bottom = 0


def add_docx_project_title(document, title, dates):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(6.2)
    table.columns[1].width = Inches(1.0)
    left = table.rows[0].cells[0].paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    right = table.rows[0].cells[1].paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    right.add_run(dates).bold = True


def build_docx():
    document = Document()
    configure_docx(document)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    run = name.add_run("ZAID SALMAN")
    run.bold = True
    run.font.name = BODY_FONT
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("GENERATIVE AI | AGENTIC AI | AI PLATFORM ENGINEER")
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    contact.add_run("Bengaluru, India | +91-8088061631 | ")
    add_hyperlink(contact, "zaid.cloudsre@gmail.com", "mailto:zaid.cloudsre@gmail.com")
    contact.add_run(" | ")
    add_hyperlink(contact, "LinkedIn", "https://linkedin.com/in/zaidsalman/")
    contact.add_run(" | ")
    add_hyperlink(contact, "GitHub", "https://github.com/Salman167")

    relocation = document.add_paragraph()
    relocation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    relocation.paragraph_format.space_after = Pt(1)
    relocation.add_run("Open to relocation across Europe, UAE, Saudi Arabia, and the wider Middle East").italic = True

    add_docx_heading(document, "Professional Summary")
    document.add_paragraph(SUMMARY)

    add_docx_heading(document, "Core Skills")
    for label, values in SKILLS:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(1.2)
        p.add_run(f"{label}: ").bold = True
        p.add_run(values)

    add_docx_heading(document, "Professional Experience")
    for role in EXPERIENCE:
        add_docx_role(document, role["title"], role["company"], role["location"], role["dates"])
        for bullet in role["bullets"]:
            add_docx_bullet(document, bullet)

    add_docx_heading(document, "Generative AI Projects")
    project = PROJECTS[0]
    add_docx_project_title(document, project["title"], project["dates"])
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    tech = p.add_run(project["tech"])
    tech.bold = True
    tech.font.size = Pt(9.2)
    tech.font.color.rgb = RGBColor.from_string(BLUE)
    for bullet in project["bullets"]:
        add_docx_bullet(document, bullet)

    document.add_page_break()
    add_docx_heading(document, "Generative AI Projects (Continued)")
    project = PROJECTS[1]
    add_docx_project_title(document, project["title"], project["dates"])
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    tech = p.add_run(project["tech"])
    tech.bold = True
    tech.font.size = Pt(9.2)
    tech.font.color.rgb = RGBColor.from_string(BLUE)
    for bullet in project["bullets"]:
        add_docx_bullet(document, bullet)

    add_docx_heading(document, "Education & Certifications")
    p = document.add_paragraph()
    p.add_run("Bachelor of Engineering in Electronics & Communication").bold = True
    p.add_run(" | SJCE College, Mysore | 2023 | CGPA: 8.53/10")
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Certifications: ").bold = True
    p.add_run(" | ".join(CERTIFICATIONS))

    document.core_properties.title = "Zaid Salman — Generative AI, Agentic AI & AI Platform Engineer"
    document.core_properties.subject = "ATS Resume"
    document.core_properties.author = "Zaid Salman"
    document.save(DOCX_PATH)


def pdf_styles():
    pdfmetrics.registerFont(TTFont(PDF_BODY_FONT, r"C:\Windows\Fonts\calibri.ttf"))
    pdfmetrics.registerFont(TTFont(PDF_BOLD_FONT, r"C:\Windows\Fonts\calibrib.ttf"))
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle(
            "Name",
            parent=styles["Normal"],
            fontName=PDF_BOLD_FONT,
            fontSize=20,
            leading=22,
            textColor=colors.HexColor(f"#{NAVY}"),
            alignment=TA_CENTER,
            spaceAfter=1,
        ),
        "title": ParagraphStyle(
            "Title",
            parent=styles["Normal"],
            fontName=PDF_BOLD_FONT,
            fontSize=11,
            leading=13,
            textColor=colors.HexColor(f"#{BLUE}"),
            alignment=TA_CENTER,
            spaceAfter=2,
        ),
        "contact": ParagraphStyle(
            "Contact",
            parent=styles["Normal"],
            fontName=PDF_BODY_FONT,
            fontSize=9,
            leading=11,
            alignment=TA_CENTER,
            spaceAfter=1,
        ),
        "section": ParagraphStyle(
            "Section",
            parent=styles["Normal"],
            fontName=PDF_BOLD_FONT,
            fontSize=11,
            leading=13,
            textColor=colors.HexColor(f"#{NAVY}"),
            borderWidth=0,
            borderPadding=0,
            spaceBefore=6,
            spaceAfter=2,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=styles["Normal"],
            fontName=PDF_BODY_FONT,
            fontSize=9.7,
            leading=11.8,
            textColor=colors.HexColor(f"#{DARK}"),
            spaceAfter=1.8,
        ),
        "role": ParagraphStyle(
            "Role",
            parent=styles["Normal"],
            fontName=PDF_BODY_FONT,
            fontSize=10,
            leading=12,
            spaceBefore=2,
            spaceAfter=0.8,
            keepWithNext=True,
        ),
        "tech": ParagraphStyle(
            "Tech",
            parent=styles["Normal"],
            fontName=PDF_BOLD_FONT,
            fontSize=8.8,
            leading=10.5,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceAfter=1.5,
            keepWithNext=True,
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            parent=styles["Normal"],
            fontName=PDF_BODY_FONT,
            fontSize=9.5,
            leading=11.6,
            leftIndent=9,
            firstLineIndent=-7,
            spaceAfter=1,
        ),
    }


def esc(text):
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("–", "&#8211;")
        .replace("—", "&#8212;")
        .replace("≈", "approximately ")
    )


def pdf_section(story, styles, title):
    story.append(Paragraph(esc(title.upper()), styles["section"]))
    story.append(
        Spacer(1, 0.1)
    )


def build_pdf():
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=11 * mm,
        bottomMargin=11 * mm,
        title="Zaid Salman — Generative AI, Agentic AI & AI Platform Engineer",
        author="Zaid Salman",
    )
    story = [
        Paragraph("ZAID SALMAN", styles["name"]),
        Paragraph("GENERATIVE AI | AGENTIC AI | AI PLATFORM ENGINEER", styles["title"]),
        Paragraph(
            'Bengaluru, India | +91-8088061631 | '
            '<link href="mailto:zaid.cloudsre@gmail.com">zaid.cloudsre@gmail.com</link> | '
            '<link href="https://linkedin.com/in/zaidsalman/">LinkedIn</link> | '
            '<link href="https://github.com/Salman167">GitHub</link>',
            styles["contact"],
        ),
        Paragraph(
            "<i>Open to relocation across Europe, UAE, Saudi Arabia, and the wider Middle East</i>",
            styles["contact"],
        ),
    ]

    pdf_section(story, styles, "Professional Summary")
    story.append(Paragraph(esc(SUMMARY), styles["body"]))

    pdf_section(story, styles, "Core Skills")
    for label, values in SKILLS:
        story.append(Paragraph(f"<b>{esc(label)}:</b> {esc(values)}", styles["body"]))

    pdf_section(story, styles, "Professional Experience")
    for role in EXPERIENCE:
        story.append(
            Paragraph(
                f"<b><font color='#{NAVY}'>{esc(role['title'])}</font></b> | "
                f"{esc(role['company'])} | {esc(role['location'])}"
                f"<font color='#{DARK}'>&nbsp;&nbsp;&nbsp;<b>{esc(role['dates'])}</b></font>",
                styles["role"],
            )
        )
        for bullet in role["bullets"]:
            story.append(Paragraph(f"&#8226;&nbsp; {esc(bullet)}", styles["bullet"]))

    pdf_section(story, styles, "Generative AI Projects")
    project = PROJECTS[0]
    story.append(
        Paragraph(
            f"<b><font color='#{NAVY}'>{esc(project['title'])}</font></b>"
            f"&nbsp;&nbsp;&nbsp;<b>{esc(project['dates'])}</b>",
            styles["role"],
        )
    )
    story.append(Paragraph(esc(project["tech"]), styles["tech"]))
    for bullet in project["bullets"]:
        story.append(Paragraph(f"&#8226;&nbsp; {esc(bullet)}", styles["bullet"]))

    story.append(PageBreak())
    pdf_section(story, styles, "Generative AI Projects (Continued)")
    project = PROJECTS[1]
    story.append(
        Paragraph(
            f"<b><font color='#{NAVY}'>{esc(project['title'])}</font></b>"
            f"&nbsp;&nbsp;&nbsp;<b>{esc(project['dates'])}</b>",
            styles["role"],
        )
    )
    story.append(Paragraph(esc(project["tech"]), styles["tech"]))
    for bullet in project["bullets"]:
        story.append(Paragraph(f"&#8226;&nbsp; {esc(bullet)}", styles["bullet"]))

    pdf_section(story, styles, "Education & Certifications")
    story.append(
        Paragraph(
            "<b>Bachelor of Engineering in Electronics &amp; Communication</b> | "
            "SJCE College, Mysore | 2023 | CGPA: 8.53/10",
            styles["body"],
        )
    )
    story.append(
        Paragraph(
            f"<b>Certifications:</b> {esc(' | '.join(CERTIFICATIONS))}",
            styles["body"],
        )
    )

    doc.build(story)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
